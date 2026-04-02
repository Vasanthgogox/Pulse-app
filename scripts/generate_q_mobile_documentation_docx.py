from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)


def add_page_number(paragraph):
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbers(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def build_document(output_path: str):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "Q Mobile End-to-End Documentation"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Q Mobile\nEnd-to-End System Documentation")
    run.bold = True
    run.font.size = Pt(20)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        f"Version 1.0 | Generated on {datetime.now().strftime('%Y-%m-%d')}"
    )
    doc.add_paragraph()

    doc.add_heading("Table of Contents", level=1)
    toc_paragraph = doc.add_paragraph(
        "(In Microsoft Word, right-click this area and choose 'Update Field' to render the TOC.)"
    )
    toc_paragraph.italic = True
    add_toc(doc.add_paragraph())
    doc.add_page_break()

    doc.add_heading("1. Overview and Purpose", level=1)
    doc.add_paragraph(
        "Q Mobile is a React Native (Expo) logistics operations application that shares the same Supabase backend as the Q-unified-base platform. "
        "It supports dispatcher/fleet workflows, finance and ledger operations, network collaboration, load marketplace operations, and a dedicated driver experience."
    )
    add_bullets(
        doc,
        [
            "Primary users: dispatchers, fleet owners, finance operators, network coordinators, and drivers.",
            "Business objective: run daily logistics operations from mobile, including trips, indents, entities, and cashflow tracking.",
            "Data platform: Supabase Auth + Postgres + RLS + RPC + optional Edge Functions for AI operations.",
            "Architecture intent: microservice-like service files by domain, thin route screens, reusable contexts and query hooks.",
        ],
    )

    doc.add_heading("2. Product Scope and Core Capabilities", level=1)
    add_bullets(
        doc,
        [
            "Authentication and session persistence with optional 'Keep me signed in'.",
            "Organization-aware data access and profile-driven capability model.",
            "Trips lifecycle management: create, assign, track, complete, and reconcile.",
            "Finance and ledger management with transaction history, filters, and entity-level views.",
            "Network hub for clients, suppliers, drivers, connection requests, and invite flows.",
            "Load board and indent creation/management.",
            "Driver application with mission flow, OTP claim, location reporting, and wallet/pay features.",
            "Ops Agent conversational UI backed by Supabase Edge Function and Gemini.",
            "Optional AI insights service for risk, prediction, and cashflow forecast datasets.",
        ],
    )

    doc.add_heading("3. High-Level Architecture", level=1)
    doc.add_heading("3.1 Technology Stack", level=2)
    add_bullets(
        doc,
        [
            "Client: Expo SDK 54, React Native 0.81, React 19, TypeScript.",
            "Navigation: Expo Router route groups for dispatcher and driver experiences.",
            "State and data: React Context + TanStack Query + realtime invalidation hooks.",
            "Backend: Supabase (Auth, Postgres, RLS, RPC, Edge Functions).",
            "Security: anon key + user JWT, RLS policies, no service_role usage in app.",
        ],
    )

    doc.add_heading("3.2 Logical Component Diagram", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "Mobile UI (app/, components/, features/) -> Context Layer (auth/org/network/language/wallet) -> Query Layer (lib/queries) -> Domain Services (features/*/services + services/*) -> Supabase (Auth, DB, RPC, Realtime, Edge Functions)"
    )

    doc.add_heading("3.3 Folder Responsibilities", level=2)
    add_bullets(
        doc,
        [
            "app/: route screens and navigation composition (thin wrappers where possible).",
            "features/: domain modules (auth, trips, finance, network, indents, drivers, vehicles, AI, ops-agent, etc.).",
            "services/: cross-feature or legacy service endpoints.",
            "contexts/: app-wide state (auth, organization, network, language, wallet, theme/avatar).",
            "lib/: shared platform utilities (supabase client, capabilities, query client, utils).",
            "docs/: operational and engineering reference documentation.",
        ],
    )

    doc.add_heading("4. End-to-End Flow", level=1)
    doc.add_heading("4.1 Application Bootstrap", level=2)
    add_numbers(
        doc,
        [
            "App starts in app/_layout.tsx and loads fonts/splash sequence.",
            "SafeAreaProvider, QueryClientProvider, and all core providers are initialized.",
            "Supabase config is validated; missing config shows a dedicated error screen.",
            "ErrorBoundary handles config, network, and session-expired states.",
            "Root index route resolves user/profile and redirects to dispatcher tabs or driver app.",
        ],
    )

    doc.add_heading("4.2 Authentication and Session Lifecycle", level=2)
    add_numbers(
        doc,
        [
            "User signs in on /sign-in with validation and online checks.",
            "Auth service executes Supabase sign-in and profile hydration.",
            "Session persists using SecureStore when possible, with AsyncStorage fallback.",
            "If Keep Signed In is disabled, app signs out on background and next cold start.",
            "Auth state subscription keeps user/profile synchronized throughout runtime.",
        ],
    )

    doc.add_heading("4.3 Dispatcher/Fleet Main Flow", level=2)
    add_numbers(
        doc,
        [
            "User lands in finance/trips/network tabs depending on navigation path.",
            "Organization context resolves current organization scope for all queries.",
            "Capabilities are derived from profile flags and gate feature visibility/actions.",
            "User executes daily operations: create/manage trips, indents, entities, and ledger transactions.",
            "Realtime invalidation updates data views when DB changes occur.",
        ],
    )

    doc.add_heading("4.4 Driver App Main Flow", level=2)
    add_numbers(
        doc,
        [
            "Driver role redirects to /(driver) route group and driver tab layout.",
            "Driver dashboard fetches assigned and OTP-pending trips.",
            "Driver accepts/claims trips, follows pickup-to-drop guidance, and updates status.",
            "Location updates are sent periodically during active missions.",
            "Trip completion triggers history/wallet consistency and mission closure.",
        ],
    )

    doc.add_heading("4.5 Textual Process Flowchart", level=2)
    doc.add_paragraph(
        "Sign-In -> Session Restore -> Role Resolution -> (Dispatcher Tabs OR Driver App) -> Domain Actions (Trips/Finance/Network/Indents) -> Supabase Writes/RPC -> Realtime Invalidation -> Updated UI -> Reports/History/Settlement"
    )

    doc.add_heading("5. Functional Modules", level=1)
    doc.add_heading("5.1 Authentication and Identity", level=2)
    add_bullets(
        doc,
        [
            "Email/password sign-in and sign-up with profile metadata.",
            "First-launch cleanup logic to prevent stale credential reuse.",
            "Session expiration recovery with controlled sign-out and route replace.",
            "Localized messaging and offline-aware login behavior.",
        ],
    )

    doc.add_heading("5.2 Organization and Permission Model", level=2)
    add_bullets(
        doc,
        [
            "Organization context fetches available organizations and sets active org.",
            "Capabilities derive from profile.role + aggregated/asset flags.",
            "Permission object controls access to trips, indents, finance, entities, and team actions.",
            "Driver role is restricted from dispatcher capabilities.",
        ],
    )

    doc.add_heading("5.3 Trips Management", level=2)
    add_bullets(
        doc,
        [
            "Create trips (manual and aggregate patterns), optional OTP generation for aggregate flow.",
            "Assignment updates for driver/vehicle with assignment audit logging.",
            "Support for client/supplier perspective views via RPC-based shared-trip retrieval.",
            "Trip statuses managed across assigned, in_progress, at_drop, completed, cancelled.",
            "Driver reject flow with backend audit and unassignment semantics.",
        ],
    )

    doc.add_heading("5.4 Finance and Ledger", level=2)
    add_bullets(
        doc,
        [
            "Organization-scoped transaction retrieval with pagination and filters.",
            "Entity-specific ledger views (party, contact ID, driver-level).",
            "Ledger entry create/update with validation and double-entry interpretation.",
            "Trip-linked finance visibility and comparison/verification UX patterns.",
            "Transaction tables integrate with realtime query invalidation for near-live updates.",
        ],
    )

    doc.add_heading("5.5 Network Hub and Load", level=2)
    add_bullets(
        doc,
        [
            "Manage network nodes (clients, suppliers, drivers) with status indicators.",
            "Send/receive/approve/reject connection requests.",
            "Invite users via in-app request or share flow fallback.",
            "Load center integration for marketplace-style indent browsing and creation.",
            "Dedicated full-screen load board route for focused operations.",
        ],
    )

    doc.add_heading("5.6 Driver Experience", level=2)
    add_bullets(
        doc,
        [
            "Dedicated driver theme/layout and custom tab bar.",
            "Trip mission guidance states: accepted, pickup, transit, reached, completed.",
            "OTP claim support for pending aggregate assignments.",
            "Map, route, and periodic location update logic during active missions.",
            "Invite/job request handling and operational feedback animations.",
        ],
    )

    doc.add_heading("5.7 Ops Agent and AI Features", level=2)
    add_bullets(
        doc,
        [
            "Ops Agent screen offers conversational operations actions and workflow prompts.",
            "Chat orchestration includes preview cards, confirmations, attachments, and PDF actions.",
            "Supabase Edge Function (ops-agent-chat) mediates Gemini calls and auth/rate controls.",
            "AI service reads optional insight tables for risk/prediction/cashflow summaries.",
        ],
    )

    doc.add_heading("6. User Guide by Persona", level=1)
    doc.add_heading("6.1 Dispatcher / Fleet Owner", level=2)
    add_numbers(
        doc,
        [
            "Sign in and ensure correct organization context is loaded.",
            "Use Trips tab to monitor active/history and create or assign new trips.",
            "Use Finance tab to record inflow/outflow and review balances by party/entity.",
            "Use Network tab to connect with clients/suppliers/drivers and process requests.",
            "Use Load board and indents for demand-supply matching and trip creation funnel.",
        ],
    )

    doc.add_heading("6.2 Finance Operator", level=2)
    add_numbers(
        doc,
        [
            "Open Finance and filter by period/entity/transaction characteristics.",
            "Post new ledger entries with correct cash-in/cash-out semantics.",
            "Validate linked trip references where applicable.",
            "Review due/paid states and reconcile with trip-level outcomes.",
        ],
    )

    doc.add_heading("6.3 Driver", level=2)
    add_numbers(
        doc,
        [
            "Sign in with driver account and open dashboard mission card.",
            "Accept assigned trip or claim pending trip via OTP.",
            "Follow pickup and drop guidance, updating status at each stage.",
            "Complete trip workflow and verify history/wallet updates.",
        ],
    )

    doc.add_heading("7. Technical Deep Dive", level=1)
    doc.add_heading("7.1 Data Access and Caching", level=2)
    add_bullets(
        doc,
        [
            "TanStack Query defaults: staleTime 60s, gcTime 5m, retry 1 for queries.",
            "Domain query hooks centralize cache keys and invalidation behavior.",
            "Realtime invalidation hooks refresh trips and transactions on DB change events.",
            "Service layer remains the source of API contracts and data transformations.",
        ],
    )

    doc.add_heading("7.2 Supabase Integration", level=2)
    add_bullets(
        doc,
        [
            "Supabase URL and anon key are read from EXPO_PUBLIC environment values.",
            "Storage strategy uses SecureStore first; AsyncStorage fallback when needed.",
            "Custom fetch adds timeout and retry for unstable network conditions.",
            "RPC endpoints support critical cross-org and assignment workflows.",
            "RLS is assumed as core security boundary for all table access.",
        ],
    )

    doc.add_heading("7.3 Routing and Navigation", level=2)
    add_bullets(
        doc,
        [
            "Root stack defines auth routes, dispatcher tabs, driver group, and modal groups.",
            "Dispatcher path uses /(tabs) routes for finance, trips, network and related screens.",
            "Driver path uses /(driver) routes with dashboard/control/history/wallet tabs.",
            "Entity detail routes (client/supplier/driver/vehicle/trip/indent) are deep-link capable.",
        ],
    )

    doc.add_heading("7.4 Safe Area and Responsive Compliance", level=2)
    add_bullets(
        doc,
        [
            "All edge-drawing screens use safe-area insets and shared layout components.",
            "Loading states should use safe-area-aware centered loading component.",
            "FAB/fixed controls account for bottom inset offset.",
            "Design guidance targets compact phones through tablets in portrait/landscape.",
        ],
    )

    doc.add_heading("8. Integrations and External Dependencies", level=1)
    add_bullets(
        doc,
        [
            "Supabase Auth, Database, Realtime, and Edge Functions.",
            "Google Gemini (via Edge Function) for Ops Agent conversational processing.",
            "Maps/location stack for route visualization and driver movement context.",
            "Expo modules (camera, contacts, secure store, sharing, print, etc.) as needed by workflows.",
        ],
    )

    doc.add_heading("9. Security, Reliability, and Governance", level=1)
    add_bullets(
        doc,
        [
            "No service-role key in mobile app; anon key + JWT only.",
            "RLS enforces user/org data boundaries.",
            "Session-expiration handling avoids stale or invalid token loops.",
            "Network retries/timeouts reduce user-facing failures on unstable connections.",
            "Edge Function deployment and secret handling are documented operationally.",
        ],
    )

    doc.add_heading("10. Deployment and Environment", level=1)
    add_numbers(
        doc,
        [
            "Install dependencies and configure EXPO_PUBLIC_SUPABASE_URL / EXPO_PUBLIC_SUPABASE_ANON_KEY.",
            "Run app via Expo start scripts for Go/dev-client/simulator scenarios.",
            "Deploy Edge Functions from repo root using Supabase CLI.",
            "Set required secrets (for example GEMINI_API_KEY) and redeploy function.",
        ],
    )

    doc.add_heading("11. Troubleshooting Guide", level=1)
    doc.add_heading("11.1 Common Issues", level=2)
    add_bullets(
        doc,
        [
            "App not configured: missing EXPO_PUBLIC Supabase variables.",
            "Network request failed: DNS/firewall/WiFi constraints; try alternate network.",
            "Session appears invalid on launch: stale refresh token; app signs out and returns to sign-in.",
            "No data visible: verify organization context and user capabilities.",
            "Function call failures in Ops Agent: verify function deployment and secrets.",
        ],
    )

    doc.add_heading("11.2 Diagnostic Checklist", level=2)
    add_numbers(
        doc,
        [
            "Confirm .env values and app restart.",
            "Validate user role/profile and active organization.",
            "Check Supabase table/RPC permissions under current user.",
            "Review Edge Function logs for chat-related errors.",
            "Retry with stable network and verify realtime updates.",
        ],
    )

    doc.add_heading("12. FAQs", level=1)
    faq = [
        ("Is this app standalone from Q-unified-base?", "No. It is a mobile client using the same Supabase backend and standards."),
        ("How are permissions controlled?", "Capabilities are derived from profile role and aggregated/asset flags, then enforced in UI and service usage."),
        ("Can drivers see dispatcher screens?", "No. Driver accounts are routed to dedicated driver routes with separate UI and flows."),
        ("How is real-time sync handled?", "TanStack Query caches data and realtime invalidation hooks refresh domain queries when backend data changes."),
        ("Where does AI processing happen?", "Ops Agent requests are sent to a Supabase Edge Function that calls Gemini; mobile app does not hold backend secrets."),
    ]
    for q, a in faq:
        q_par = doc.add_paragraph()
        q_run = q_par.add_run(f"Q: {q}")
        q_run.bold = True
        doc.add_paragraph(f"A: {a}")

    doc.add_heading("13. Appendices", level=1)
    doc.add_heading("13.1 Glossary", level=2)
    glossary = [
        ("RLS", "Row Level Security in Supabase/Postgres."),
        ("Org", "Organization scope for multi-tenant data partitioning."),
        ("Aggregate Trip", "Trip flow involving supplier-side or network-linked assignment patterns."),
        ("Indent", "Load request/demand record used for matching and execution workflows."),
        ("Ops Agent", "Conversational assistant interface for operational actions."),
    ]
    for term, desc in glossary:
        p = doc.add_paragraph()
        r = p.add_run(f"{term}: ")
        r.bold = True
        p.add_run(desc)

    doc.add_heading("13.2 Domain Services Inventory", level=2)
    add_bullets(
        doc,
        [
            "auth.service.ts, organization.service.ts",
            "trips.service.ts, tripOtp.service.ts, trip-assignment-audit.service.ts",
            "finance.service.ts, indents.service.ts, direct-quotes.service.ts",
            "clients.service.ts, suppliers.service.ts, drivers.service.ts, vehicles.service.ts",
            "connectionRequestsService.ts, sharedLedgerService.ts, tripDocumentsService.ts",
            "ai.service.ts and ops-agent Edge Function integration",
        ],
    )

    doc.add_heading("13.3 Operational Readiness Checklist", level=2)
    add_numbers(
        doc,
        [
            "Environment variables configured for target environment.",
            "Supabase project linked and key RPC functions available.",
            "Role/capability expectations validated for each persona.",
            "Edge function deployed and secrets configured.",
            "Critical user journeys smoke-tested on iOS and Android.",
        ],
    )

    doc.add_paragraph()
    end_note = doc.add_paragraph(
        "End of Document"
    )
    end_note.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(output_path)


if __name__ == "__main__":
    OUTPUT = "docs/Q-Mobile_End_to_End_Documentation.docx"
    build_document(OUTPUT)
    print(f"Generated: {OUTPUT}")
